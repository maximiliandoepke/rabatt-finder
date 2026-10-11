"""Erzeugt die Excel-Übersicht (Suche, Übersicht, Top-Angebote, Anbieter-Vergleich, Alle Angebote)."""
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .parse import CB, PLAT_KURZ, PLATTS, SB, UD


def build_excel(rows, out, stand):
    STAND = stand
    OUT = out
    N = len(rows)
    LAST = N + 1
    KATEGORIEN = sorted({r["Kategorie"] for r in rows})
    counts = {p: sum(r["Plattform"] == p for r in rows) for p in PLATTS}

    # ---------- Excel ----------
    BLUE = "1F4E78"
    HEAD_FILL = PatternFill("solid", fgColor=BLUE)
    HEAD_FONT = Font(bold=True, color="FFFFFF")
    SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
    TOP_FILL = PatternFill("solid", fgColor="C6EFCE")
    TOP_FONT = Font(bold=True, color="006100")
    HIT_FILL = PatternFill("solid", fgColor="FFF2CC")
    LINK_FONT = Font(color="0563C1", underline="single")
    GREY = Font(italic=True, color="595959")
    thin = Side(style="thin", color="BFBFBF")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    TOP_ALIGN = Alignment(vertical="top", wrap_text=True)


    def style_header(cells):
        for c in cells:
            c.fill, c.font = HEAD_FILL, HEAD_FONT
            c.alignment = Alignment(vertical="center", wrap_text=True)


    def widths(ws, ws_widths):
        for col, w in ws_widths.items():
            ws.column_dimensions[col].width = w


    def title(ws, text, sub=None):
        ws["A1"] = text
        ws["A1"].font = Font(bold=True, size=16, color=BLUE)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = GREY


    def back_link(ws, cell="A3"):
        ws[cell] = "← zurück zum Start"
        ws[cell].hyperlink = "#'Start'!A1"
        ws[cell].font = LINK_FONT


    def pct_cell(c, v):
        c.value = v
        c.number_format = "0%"


    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    start = wb.active
    start.title = "Start"
    su = wb.create_sheet("Suche")
    ov = wb.create_sheet("Übersicht")
    top = wb.create_sheet("Top-Angebote")
    cmp_ws = wb.create_sheet("Anbieter-Vergleich")
    ws = wb.create_sheet("Alle Angebote")
    li = wb.create_sheet("Listen")
    info = wb.create_sheet("Info")

    # ---------- Alle Angebote ----------
    COLS = ["Plattform", "Anbieter", "Angebot", "Kategorie", "Max. Rabatt", "Rabatt-Art", "Kanal", "Unterkategorie",
            "Beschreibung", "Hinweis", "Schlagworte", "Link", "Treffer", "Sortierschlüssel"]
    # A..N
    ws.append(COLS)
    for i, row in enumerate(rows):
        r = i + 2
        ws.append([row[c] for c in COLS[:12]])
        ws[f"E{r}"].number_format = "0%"
        ws[f"L{r}"].hyperlink = row["Link"]
        ws[f"L{r}"].font = LINK_FONT
        ws[f"M{r}"] = (f'=IF(AND(OR(TRIM(Suche!$C$4)="",ISNUMBER(SEARCH(TRIM(Suche!$C$4),'
                       f'B{r}&" "&C{r}&" "&D{r}&" "&H{r}&" "&I{r}&" "&J{r}&" "&K{r}))),'
                       f'OR(Suche!$C$5="",Suche!$C$5="Alle",D{r}=Suche!$C$5),'
                       f'OR(Suche!$C$6="",Suche!$C$6="Alle",A{r}=Suche!$C$6),'
                       f'OR(Suche!$C$7="",Suche!$C$7="Alle",ISNUMBER(SEARCH(Suche!$C$7,G{r})))),1,0)')
        ws[f"N{r}"] = f'=IF(M{r}=1,N(E{r})+(100000-ROW())/100000000,"")'
    style_header(ws[1])
    widths(ws, {"A": 24, "B": 28, "C": 46, "D": 26, "E": 11, "F": 15, "G": 15, "H": 30, "I": 60, "J": 30, "K": 40,
                "L": 40, "M": 8, "N": 8})
    ws.column_dimensions["N"].hidden = True
    ws.freeze_panes = "C2"
    t = Table(displayName="AlleAngebote", ref=f"A1:N{LAST}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.conditional_formatting.add(f"A2:K{LAST}", FormulaRule(formula=['AND(TRIM(Suche!$C$4)<>"",$M2=1)'], fill=HIT_FILL))

    # ---------- Listen (Dropdowns) ----------
    lists = {"A": ["Alle"] + KATEGORIEN, "B": ["Alle"] + PLATTS, "C": ["Alle", "online", "vor Ort"]}
    for col, vals in lists.items():
        for i, v in enumerate(vals, 1):
            li[f"{col}{i}"] = v
    li.sheet_state = "hidden"

    # ---------- Suche ----------
    AL = "'Alle Angebote'"
    RNG = lambda col: f"{AL}!${col}$2:${col}${LAST}"  # noqa: E731
    title(su, "Angebotssuche",
          "Suchbegriff eingeben (z. B. Fernseher, Hotel, Laptop, Schuhe) und/oder Filter wählen. Leeres Suchfeld = alle Angebote.")
    back_link(su, "F1")
    filters = [("Suchbegriff:", "", None), ("Kategorie:", "Alle", "A"), ("Plattform:", "Alle", "B"), ("Kanal:", "Alle", "C")]
    for r, (label, val, lcol) in enumerate(filters, start=4):
        su[f"A{r}"] = label
        su[f"A{r}"].font = Font(bold=True)
        su[f"C{r}"] = val
        su[f"C{r}"].fill, su[f"C{r}"].border = HIT_FILL, BOX
        su[f"C{r}"].font = Font(bold=True, size=12)
        su.merge_cells(f"C{r}:D{r}")
        if lcol:
            dv = DataValidation(type="list", formula1=f"=Listen!${lcol}$1:${lcol}${len(lists[lcol])}", allow_blank=True)
            su.add_data_validation(dv)
            dv.add(f"C{r}")
    su["E4"] = "Tipp: Wortteile reichen, z. B. „fernseh“ oder „hotel“."
    su["E4"].font = GREY

    su["A9"] = "Treffer gesamt:"
    su["A9"].font = Font(bold=True)
    su["C9"] = f"=COUNTIF({RNG('M')},1)"
    su["C9"].font = Font(bold=True, size=12)

    su["A11"], su["C11"], su["D11"] = "Plattform", "Treffer", "Bester Rabatt"
    style_header([su["A11"], su["B11"], su["C11"], su["D11"]])
    su.merge_cells("A11:B11")
    for i, p in enumerate(PLATTS):
        r = 12 + i
        su[f"A{r}"] = p
        su.merge_cells(f"A{r}:B{r}")
        su[f"C{r}"] = f"=COUNTIFS({RNG('A')},A{r},{RNG('M')},1)"
        su[f"D{r}"] = f"=SUMPRODUCT(MAX(({RNG('A')}=A{r})*({RNG('M')}=1)*{RNG('E')}))"
        su[f"D{r}"].number_format = "0%"
        for col in "ABCD":
            su[f"{col}{r}"].border = BOX
    su.conditional_formatting.add(
        "A12:D14", FormulaRule(formula=["AND($D12>0,$D12=MAX($D$12:$D$14))"], fill=TOP_FILL, font=TOP_FONT))

    HR = 17
    su["F11"] = "Bestes Angebot:"
    su["F11"].font = Font(bold=True)
    su["F12"] = f'=IF(A{HR + 1}="","Keine Treffer",B{HR + 1}&" ("&A{HR + 1}&")")'
    su["F13"] = f'=IF(E{HR + 1}="","",TEXT(E{HR + 1},"0%")&" – "&C{HR + 1})'
    su["F12"].font = Font(bold=True, size=12, color="006100")

    heads = ["Plattform", "Anbieter", "Angebot", "Kategorie", "Max. Rabatt", "Kanal", "Hinweis", "Link"]
    src_cols = {1: "A", 2: "B", 3: "C", 4: "D", 6: "G", 7: "J"}
    for j, h in enumerate(heads, 1):
        su.cell(row=HR, column=j, value=h)
    style_header(su[HR][:len(heads)])
    for k in range(1, N + 1):
        r = HR + k
        su[f"P{r}"] = f'=IFERROR(LARGE({RNG("N")},{k}),"")'
        su[f"Q{r}"] = f'=IF(P{r}="","",MATCH(P{r},{RNG("N")},0))'
        for col_idx, src in src_cols.items():
            su.cell(row=r, column=col_idx, value=f'=IF($Q{r}="","",INDEX({RNG(src)},$Q{r})&"")')
        su[f"E{r}"] = f'=IF($Q{r}="","",IF(INDEX({RNG("E")},$Q{r})="","",INDEX({RNG("E")},$Q{r})))'
        su[f"E{r}"].number_format = "0%"
        su[f"H{r}"] = f'=IF($Q{r}="","",HYPERLINK(INDEX({RNG("L")},$Q{r}),"öffnen"))'
        su[f"H{r}"].font = LINK_FONT
    su.conditional_formatting.add(
        f"A{HR + 1}:G{HR + N}",
        FormulaRule(formula=[f'AND($E{HR + 1}<>"",N($E{HR + 1})>0,$E{HR + 1}=MAX($E${HR + 1}:$E${HR + N}))'],
                    fill=TOP_FILL, font=TOP_FONT))
    su.column_dimensions["P"].hidden = True
    su.column_dimensions["Q"].hidden = True
    widths(su, {"A": 26, "B": 30, "C": 50, "D": 26, "E": 12, "F": 18, "G": 34, "H": 10})
    su.freeze_panes = f"A{HR + 1}"


    # ---------- Hilfsfunktionen für Auswertungen ----------
    def best(items):
        with_pct = [r for r in items if r["Max. Rabatt"] is not None]
        return max(with_pct, key=lambda r: r["Max. Rabatt"]) if with_pct else None


    # ---------- Top-Angebote ----------
    title(top, "Top-Angebote je Kategorie",
          "Die 5 höchsten Rabatte je Kategorie (pro Anbieter nur das beste Angebot). „Bis zu“-Rabatte gelten oft nur für einzelne Artikel.")
    back_link(top)
    top_anchor = {}
    r = 5
    for kat in KATEGORIEN:
        top_anchor[kat] = r
        top[f"A{r}"] = kat
        top[f"A{r}"].font = Font(bold=True, size=12, color=BLUE)
        for col in "ABCDEFGH":
            top[f"{col}{r}"].fill = SEC_FILL
        r += 1
        for j, h in enumerate(["Rang", "Plattform", "Anbieter", "Angebot", "Max. Rabatt", "Kanal", "Link"], 2):
            c = top.cell(row=r, column=j, value=h)
            c.font = Font(bold=True, color=BLUE)
        r += 1
        seen_a, rank = set(), 0
        for row in sorted((x for x in rows if x["Kategorie"] == kat and x["Max. Rabatt"] is not None),
                          key=lambda x: (-x["Max. Rabatt"], PLATTS.index(x["Plattform"]))):
            akey = (row["Anbieter"].lower(), row["Plattform"])
            if akey in seen_a:
                continue
            seen_a.add(akey)
            rank += 1
            top[f"B{r}"] = rank
            top[f"C{r}"] = PLAT_KURZ[row["Plattform"]]
            top[f"D{r}"] = row["Anbieter"]
            top[f"E{r}"] = row["Angebot"]
            pct_cell(top[f"F{r}"], row["Max. Rabatt"])
            top[f"G{r}"] = row["Kanal"]
            top[f"H{r}"] = "öffnen"
            top[f"H{r}"].hyperlink = row["Link"]
            top[f"H{r}"].font = LINK_FONT
            if rank == 1:
                for col in "BCDEFG":
                    top[f"{col}{r}"].fill = TOP_FILL
                top[f"D{r}"].font = TOP_FONT
            for col in "BCDEFGH":
                top[f"{col}{r}"].alignment = TOP_ALIGN
            r += 1
            if rank == 5:
                break
        r += 1
    widths(top, {"A": 4, "B": 6, "C": 20, "D": 30, "E": 60, "F": 12, "G": 16, "H": 10})

    # ---------- Übersicht ----------
    title(ov, "Übersicht nach Kategorie",
          "Anzahl Angebote und bester Rabatt je Plattform. Grün = Plattform mit dem höchsten Rabatt in der Kategorie. "
          "Klick auf eine Kategorie springt zu deren Top-Angeboten.")
    back_link(ov)
    OH = 5
    ov_heads = ["Kategorie"]
    for p in PLATTS:
        ov_heads += [f"{PLAT_KURZ[p]}\nAnzahl", f"{PLAT_KURZ[p]}\nbester Rabatt"]
    ov_heads += ["Gesamt", "Beste Plattform", "Bestes Angebot"]
    for j, h in enumerate(ov_heads, 1):
        ov.cell(row=OH, column=j, value=h)
    style_header(ov[OH][:len(ov_heads)])
    ov.row_dimensions[OH].height = 32
    r = OH + 1
    for kat in KATEGORIEN:
        items = [x for x in rows if x["Kategorie"] == kat]
        c = ov.cell(row=r, column=1, value=kat)
        c.hyperlink = f"#'Top-Angebote'!A{top_anchor[kat]}"
        c.font = LINK_FONT
        bests = {}
        for i, p in enumerate(PLATTS):
            pit = [x for x in items if x["Plattform"] == p]
            ov.cell(row=r, column=2 + 2 * i, value=len(pit) or None)
            b = best(pit)
            bests[p] = b
            if b:
                pct_cell(ov.cell(row=r, column=3 + 2 * i), b["Max. Rabatt"])
        ov.cell(row=r, column=8, value=len(items))
        overall = best([b for b in bests.values() if b])
        if overall:
            ov.cell(row=r, column=9, value=PLAT_KURZ[overall["Plattform"]])
            ov.cell(row=r, column=10, value=f'{overall["Anbieter"]} – {overall["Angebot"]}')
            top_pct = overall["Max. Rabatt"]
            for i, p in enumerate(PLATTS):
                if bests[p] and bests[p]["Max. Rabatt"] == top_pct:
                    for col in (2 + 2 * i, 3 + 2 * i):
                        ov.cell(row=r, column=col).fill = TOP_FILL
                        ov.cell(row=r, column=col).font = TOP_FONT
        for col in range(1, 11):
            ov.cell(row=r, column=col).border = BOX
            ov.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=col == 10)
        r += 1
    ov.cell(row=r, column=1, value="Gesamt").font = Font(bold=True)
    for i, p in enumerate(PLATTS):
        ov.cell(row=r, column=2 + 2 * i, value=counts[p]).font = Font(bold=True)
    ov.cell(row=r, column=8, value=N).font = Font(bold=True)
    for col in range(1, 11):
        ov.cell(row=r, column=col).border = BOX
        ov.cell(row=r, column=col).fill = SEC_FILL

    r += 3
    ov.cell(row=r, column=1, value="Angebote nach Rabatt-Art und Kanal").font = Font(bold=True, size=12, color=BLUE)
    r += 1
    arten = ["Prozent", "Betrag", "Gratis / Test", "Aktion / Sonstiges"]
    kanäle = ["online", "online & vor Ort", "vor Ort", "nur vor Ort"]
    hdr = ["", *[PLAT_KURZ[p] for p in PLATTS], "Gesamt"]
    for j, h in enumerate(hdr, 1):
        ov.cell(row=r, column=j, value=h)
    style_header(ov[r][:len(hdr)])
    r += 1
    for label, key, vals in [("Rabatt-Art", "Rabatt-Art", arten), ("Kanal", "Kanal", kanäle)]:
        ov.cell(row=r, column=1, value=label).font = Font(bold=True, color=BLUE)
        r += 1
        for v in vals:
            n_all = sum(x[key] == v for x in rows)
            if not n_all:
                continue
            ov.cell(row=r, column=1, value="   " + v)
            for i, p in enumerate(PLATTS):
                ov.cell(row=r, column=2 + i, value=sum(x[key] == v and x["Plattform"] == p for x in rows) or None)
            ov.cell(row=r, column=5, value=n_all)
            r += 1
    widths(ov, {"A": 30, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 10, "I": 20, "J": 70})
    ov.freeze_panes = f"B{OH + 1}"


    # ---------- Anbieter-Vergleich ----------
    def norm(name):
        s = name.lower()
        s = re.sub(r"[®™©’'`´]", "", s)
        s = re.sub(r"\b(geschenkgutscheine?|gutscheine?|gift ?cards?|online ?shop|deutschland|germany|international|"
                   r"store|shop|de)\b", "", s)
        return re.sub(r"[^a-z0-9äöüß]+", "", s)


    groups = {}
    for row in rows:
        k = norm(row["Anbieter"])
        if len(k) >= 2:
            groups.setdefault(k, []).append(row)
    multi = {k: v for k, v in groups.items() if len({x["Plattform"] for x in v}) >= 2}

    title(cmp_ws, "Anbieter-Vergleich",
          f"{len(multi)} Anbieter gibt es auf mehreren Plattformen. Grün = Plattform mit dem höchsten Rabatt; "
          "„ohne %“ = Angebot vorhanden, aber ohne Prozentrabatt.")
    back_link(cmp_ws)
    CH = 5
    c_heads = ["Anbieter", "Kategorie", *[PLAT_KURZ[p] for p in PLATTS], "Beste Plattform", "Bestes Angebot", "Link"]
    for j, h in enumerate(c_heads, 1):
        cmp_ws.cell(row=CH, column=j, value=h)
    style_header(cmp_ws[CH][:len(c_heads)])
    r = CH + 1
    for k, items in sorted(multi.items(), key=lambda kv: kv[1][0]["Anbieter"].lower()):
        name = min((x["Anbieter"] for x in items), key=len)
        kat = max({x["Kategorie"] for x in items}, key=lambda c: sum(y["Kategorie"] == c for y in items))
        cmp_ws.cell(row=r, column=1, value=name)
        cmp_ws.cell(row=r, column=2, value=kat)
        bests = {}
        for i, p in enumerate(PLATTS):
            pit = [x for x in items if x["Plattform"] == p]
            c = cmp_ws.cell(row=r, column=3 + i)
            b = best(pit)
            bests[p] = b
            if b:
                pct_cell(c, b["Max. Rabatt"])
            elif pit:
                c.value = "ohne %"
                c.font = GREY
            else:
                c.value = "–"
                c.font = GREY
            c.alignment = Alignment(horizontal="center")
        overall = best([b for b in bests.values() if b])
        if overall:
            cmp_ws.cell(row=r, column=6, value=PLAT_KURZ[overall["Plattform"]]).font = TOP_FONT
            cmp_ws.cell(row=r, column=7, value=overall["Angebot"])
            lc = cmp_ws.cell(row=r, column=8, value="öffnen")
            lc.hyperlink = overall["Link"]
            lc.font = LINK_FONT
            for i, p in enumerate(PLATTS):
                if bests[p] and bests[p]["Max. Rabatt"] == overall["Max. Rabatt"]:
                    cmp_ws.cell(row=r, column=3 + i).fill = TOP_FILL
                    cmp_ws.cell(row=r, column=3 + i).font = TOP_FONT
        else:
            cmp_ws.cell(row=r, column=6, value="kein %-Rabatt").font = GREY
        for col in range(1, 9):
            cmp_ws.cell(row=r, column=col).border = BOX
        r += 1
    cmp_ws.auto_filter.ref = f"A{CH}:H{r - 1}"
    widths(cmp_ws, {"A": 28, "B": 26, "C": 18, "D": 15, "E": 12, "F": 18, "G": 60, "H": 10})
    cmp_ws.freeze_panes = f"B{CH + 1}"

    # ---------- Start ----------
    title(start, "Mitarbeiter- & Studi-Rabatte", f"Stand {STAND} · {N:,} Angebote".replace(",", "."))
    r = 4
    for p in PLATTS:
        start[f"A{r}"] = PLAT_KURZ[p]
        start[f"B{r}"] = counts[p]
        start[f"B{r}"].alignment = Alignment(horizontal="left")
        r += 1
    r += 1
    start[f"A{r}"] = "Blätter"
    start[f"A{r}"].font = Font(bold=True, size=12, color=BLUE)
    r += 1
    sheets = [
        ("Suche", "Begriff eingeben (z. B. „Fernseher“) und nach Kategorie, Plattform und Kanal filtern – "
                  "Ergebnisse nach Rabatt sortiert, bester Rabatt grün."),
        ("Übersicht", "Alle Kategorien auf einen Blick: Anzahl Angebote und bester Rabatt je Plattform."),
        ("Top-Angebote", "Die 5 besten Angebote je Kategorie."),
        ("Anbieter-Vergleich", f"{len(multi)} Marken, die es auf mehreren Plattformen gibt – wo ist der Rabatt am höchsten?"),
        ("Alle Angebote", "Die komplette Liste mit Filtern in jeder Spalte."),
        ("Info", "Quellen, Stand und Erklärungen zu den Spalten."),
    ]
    for name, desc in sheets:
        c = start[f"A{r}"]
        c.value = name
        c.hyperlink = f"#'{name}'!A1"
        c.font = Font(color="0563C1", underline="single", bold=True)
        start[f"B{r}"] = desc
        start[f"B{r}"].alignment = TOP_ALIGN
        r += 1
    r += 1
    start[f"A{r}"] = "So geht's"
    start[f"A{r}"].font = Font(bold=True, size=12, color=BLUE)
    r += 1
    for tip in [
        "Etwas Bestimmtes suchen → Blatt „Suche“, Begriff in das gelbe Feld eintragen.",
        "Stöbern nach Themen → Blatt „Übersicht“, auf eine Kategorie klicken.",
        "Gleiche Marke, wo am günstigsten? → Blatt „Anbieter-Vergleich“.",
        "Nur Angebote vor Ort/in Filialen → in der Suche „Kanal: vor Ort“ wählen.",
        "Jeder Eintrag hat einen Link „öffnen“, der direkt zum Angebot bzw. zur Kategorie führt (Login erforderlich).",
    ]:
        start[f"A{r}"] = "•"
        start[f"A{r}"].alignment = Alignment(horizontal="right", vertical="top")
        start[f"B{r}"] = tip
        start[f"B{r}"].alignment = TOP_ALIGN
        r += 1
    widths(start, {"A": 22, "B": 110})

    # ---------- Info ----------
    for line in [
        ("Stand", STAND),
        ("Angebote", f"{N} gesamt – " + ", ".join(f"{PLAT_KURZ[p]}: {n}" for p, n in counts.items())),
        ("Quelle 1", "miele.mitarbeiterangebote.de – alle 13 Kategorien (inkl. Regionales)"),
        ("Quelle 2", "studentbeans.com/de – alle 11 Kategorien, vollständig nachgeladen"),
        ("Quelle 3", "myunidays.com/DE – alle Kategorien und Unterkategorien, vollständig nachgeladen"),
        ("Kategorie", "Einheitliche Kategorie über alle Plattformen; die Original-Rubrik steht in „Unterkategorie“."),
        ("Max. Rabatt", "Höchster genannter Prozentwert. „Bis zu“-Rabatte gelten oft nur für einzelne Artikel."),
        ("Rabatt-Art", "Prozent / Betrag (z. B. 10 €) / Gratis / Test (z. B. Probemonate) / Aktion / Sonstiges."),
        ("Kanal", "online · online & vor Ort (Filialen vorhanden) · vor Ort (im Geschäft) · nur vor Ort (regionale Angebote)."),
        ("Link", "Corporate Benefits und Student Beans: direkt zum Angebot. UNiDAYS: zur Kategorieseite, auf der das Angebot steht."),
        ("Schlagworte", "Automatisch aus den Unterkategorien abgeleitet; eigene Begriffe können in „Alle Angebote“ ergänzt werden."),
        ("Anbieter-Vergleich", "Anbieter werden über den Namen zugeordnet; abweichende Schreibweisen können fehlen."),
    ]:
        info.append(line)
    widths(info, {"A": 20, "B": 120})
    for c in info["A"]:
        c.font = Font(bold=True)
    back_link(info, "D1")

    wb.save(OUT)
    return {"anzahl": N, "plattformen": counts, "vergleich": len(multi)}
